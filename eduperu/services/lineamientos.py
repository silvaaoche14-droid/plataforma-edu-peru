"""Lineamientos por curso: traerlos del MINEDU (con IA) o leer los que sube el docente (PDF / Word).
Se guardan SOLO para ese curso (no cambian el currículo de los demás) y siempre se verifican contra el texto."""
import io
import re

from eduperu.services.curriculo import competencias_de
from eduperu.core.db import execute, execute_many, fetch
from eduperu.services.programas import (buscar_programas, descargar_y_leer, extraer_items, nivel_probable, paginas_guardadas,
                       poner_en_cola)
from eduperu.core.repo import actualizar_asignacion
from eduperu.core.util import norm


def paginas_de_archivo(nombre, datos: bytes):
    """Convierte un PDF, Word o TXT en páginas [(n, texto, normalizado)]."""
    bajo = nombre.lower()
    if bajo.endswith(".pdf"):
        from pypdf import PdfReader
        lector = PdfReader(io.BytesIO(datos))
        textos = [(p.extract_text() or "") for p in lector.pages]
    elif bajo.endswith(".docx"):
        from docx import Document
        d = Document(io.BytesIO(datos))
        lineas = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for fila in t.rows:
                lineas.append(" | ".join(c.text.strip() for c in fila.cells))
        textos, actual = [], ""
        for l in lineas:                         # páginas artificiales de ~3000 caracteres
            if len(actual) + len(l) > 3000:
                textos.append(actual)
                actual = ""
            actual += l + "\n"
        textos.append(actual)
    elif bajo.endswith(".txt"):
        crudo = datos.decode("utf-8", errors="ignore")
        textos = [crudo[i:i + 3000] for i in range(0, len(crudo), 3000)] or [""]
    else:
        raise RuntimeError("Solo se pueden leer archivos PDF, Word (.docx) o TXT.")
    return [(i, t, norm(t)) for i, t in enumerate(textos, 1)]


def _guardar_propios(asig, items, fuente):
    """Guarda lo extraído como lineamientos propios del curso (reemplaza lo propio anterior de esas competencias)."""
    aid = asig["id"]
    for cid in {i[0] for i in items}:
        execute("DELETE FROM curriculo_desempenos WHERE competencia_id=%s AND asignacion_id=%s", (cid, aid))
        execute("DELETE FROM curriculo_estandares WHERE competencia_id=%s AND asignacion_id=%s", (cid, aid))
    des = [(i[0], i[2], i[3], k, aid) for k, i in enumerate((x for x in items if x[1] == "desempeno"), 1)]
    est = [(i[0], i[2], i[3], aid) for i in items if i[1] == "estandar"]
    execute_many("INSERT INTO curriculo_desempenos (competencia_id, grado, texto, orden, asignacion_id) VALUES (%s,%s,%s,%s,%s)", des)
    execute_many("INSERT INTO curriculo_estandares (competencia_id, ciclo, texto, asignacion_id) VALUES (%s,%s,%s,%s)", est)
    actualizar_asignacion(aid, lineamientos_fuente=fuente[:300])
    return len(des), len(est)


def _trabajos(asig):
    return [(asig["area"], c) for c in competencias_de(asig["nivel"], asig["area"])]


def desde_archivo(asig, nombre, datos, progreso=None):
    """El docente sube sus lineamientos: se lee, se extrae con IA y se verifica contra su propio archivo."""
    paginas = paginas_de_archivo(nombre, datos)
    items, resumen = extraer_items(paginas, asig["nivel"], _trabajos(asig), progreso)
    if not items:
        raise RuntimeError("No se encontraron desempeños en tu archivo para las competencias de este curso. "
                           "Revisa que el archivo contenga los desempeños por grado.")
    resumen["desempenos"], resumen["estandares"] = _guardar_propios(asig, items, f"Archivo propio: {nombre}")
    return resumen


def _pdf_del_nivel(nivel):
    filas = fetch("SELECT * FROM curriculo_pdf WHERE estado IN ('leido','actualizado') AND paginas>0 "
                  "AND nivel=%s ORDER BY (estado='leido') DESC, id DESC", (nivel,))
    return filas[0] if filas else None


def traer_oficial(asig, progreso=None):
    """Busca el Programa Curricular oficial del nivel (si no está leído lo descarga) y trae los lineamientos del curso."""
    pdf = _pdf_del_nivel(asig["nivel"])
    if pdf is None:
        if progreso:
            progreso(0.05, "Buscando el Programa Curricular en gob.pe…")
        lista, _ = buscar_programas()
        cand = [p for p in lista if p["nivel"] == asig["nivel"] and p["estado"] != "descartado"]
        if not cand:
            raise RuntimeError("No encontré el Programa Curricular de tu nivel en las páginas oficiales. "
                               "Pídele al administrador que lo cargue, o sube tus lineamientos en PDF o Word.")
        if progreso:
            progreso(0.1, "Descargando el PDF oficial (puede tardar)…")
        pdf = cand[0]
        descargar_y_leer(pdf["id"])
    paginas = paginas_guardadas(pdf["id"])
    items, resumen = extraer_items(paginas, asig["nivel"], _trabajos(asig),
                                   (lambda f, t: progreso(0.2 + 0.8 * f, t)) if progreso else None)
    if not items:
        raise RuntimeError("No pude verificar desempeños de este curso en el PDF oficial. "
                           "Puedes subir tus lineamientos en PDF o Word.")
    poner_en_cola(pdf["id"], items)        # el administrador podrá aprobarlos para todos
    resumen["desempenos"], resumen["estandares"] = _guardar_propios(asig, items, f"MINEDU (PDF oficial): {pdf['titulo'] or pdf['url']}")
    return resumen


def usar_oficial(asig):
    """Vuelve a usar el currículo oficial del sistema (borra los lineamientos propios del curso)."""
    execute("DELETE FROM curriculo_desempenos WHERE asignacion_id=%s", (asig["id"],))
    execute("DELETE FROM curriculo_estandares WHERE asignacion_id=%s", (asig["id"],))
    actualizar_asignacion(asig["id"], lineamientos_fuente="")
