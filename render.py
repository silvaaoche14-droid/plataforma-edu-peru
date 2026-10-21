"""Documentos oficiales: se arman como "bloques" y se dibujan igual en Word y en la vista previa."""
import io
import html as _html
from typing import Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from util import s

COLORES = {"Azul": "1F3B8F", "Morado": "6D28D9", "Verde": "15803D", "Rojo": "B91C1C",
           "Marrón": "7C4A2D", "Naranja": "C2410C", "Gris": "475569"}
ICONO_COLOR = {"Azul": "🔵", "Morado": "🟣", "Verde": "🟢", "Rojo": "🔴", "Marrón": "🟤", "Naranja": "🟠", "Gris": "⚪"}
COLOR_DEFECTO = "1F3B8F"


def _bul(items) -> str:
    """Lista → texto con viñetas (una por línea)."""
    if not items:
        return ""
    if isinstance(items, str):
        items = [x for x in items.split("\n") if x.strip()]
    return "\n".join("• " + s(x) for x in items if s(x))


def _aclarar(hex_color: str, factor: float = 0.88) -> str:
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    r, g, b = (int(c + (255 - c) * factor) for c in (r, g, b))
    return f"{r:02X}{g:02X}{b:02X}"


# ------------------------------------------------------------------
# CONSTRUCTORES DE BLOQUES
# Bloques: ("h", titulo) ("sec", texto) ("kv", [(k, v)], ncols) ("tabla", encabezados, filas)
#          ("p", texto) ("sub", texto) ("ul", [items]) ("salto",)
# ------------------------------------------------------------------
def _kv_datos(asig, extra=None):
    pares = [("Institución educativa", asig.get("institucion")), ("Director(a)", asig.get("director")),
             ("Docente", asig.get("docente")), ("Región", asig.get("region")),
             ("Nivel", asig.get("nivel")), ("Grado / edad", asig.get("grado")),
             ("Área curricular", asig.get("area")), ("Sección", asig.get("seccion")),
             ("Ciclo escolar", asig.get("ciclo_escolar"))]
    return pares + list(extra or [])


def _bloques_transversales(trans, enfoques_sel=None):
    comps_t, enfoques_db = trans
    b = [("sec", "COMPETENCIAS TRANSVERSALES"),
         ("tabla", ["Competencia transversal", "Capacidades"],
          [[c["nombre"], _bul(c["items"])] for c in comps_t])]
    if enfoques_sel:
        b += [("sec", "ENFOQUES TRANSVERSALES"),
              ("tabla", ["Enfoque transversal", "Valores", "Acciones observables"],
               [[e.get("enfoque"), e.get("valores"), e.get("acciones")] for e in enfoques_sel])]
    else:
        b += [("sec", "ENFOQUES TRANSVERSALES"),
              ("tabla", ["Enfoque transversal", "Valores"],
               [[e["nombre"], _bul(e["items"])] for e in enfoques_db])]
    return b


def bloques_anual(asig, d, comps, trans):
    b = [("h", f"PLANIFICACIÓN ANUAL {asig.get('ciclo_escolar', '')} — {asig.get('area', '')}"),
         ("sec", "I. DATOS INFORMATIVOS"), ("kv", _kv_datos(asig), 2),
         ("sec", "II. CARACTERIZACIÓN Y CONTEXTO DEL ESTUDIANTE"), ("p", asig.get("caracterizacion") or ""),
         ("sec", "III. COMPETENCIAS Y CAPACIDADES DEL ÁREA"),
         ("tabla", ["Competencia", "Capacidades"], [[c["nombre"], _bul(c["capacidades"])] for c in comps]),
         ("sec", "IV. SITUACIONES SIGNIFICATIVAS Y UNIDADES DE APRENDIZAJE")]
    filas = [[f"Unidad {u.get('numero')}", u.get("titulo"), u.get("problema"), u.get("situacion"), u.get("valores"),
              _bul(u.get("competencias")), _bul(u.get("capacidades")), f"{u.get('semanas', '')} sem."]
             for u in d.get("unidades", [])]
    b.append(("tabla", ["Unidad", "Título", "Problema identificado", "Situación significativa",
                        "Valores observables", "Competencias", "Capacidades a trabajar", "Duración"], filas))
    return b + _bloques_transversales(trans)


def bloques_unidad(asig, d, trans):
    b = [("h", f"UNIDAD DE APRENDIZAJE N.° {d.get('numero', '')}: {d.get('titulo', '')}"),
         ("sec", "I. DATOS INFORMATIVOS"),
         ("kv", _kv_datos(asig, [("Unidad", d.get("numero")), ("Duración", f"{d.get('semanas', '')} semanas")]), 2),
         ("sec", "II. SITUACIÓN SIGNIFICATIVA"), ("p", d.get("situacion")),
         ("sec", "III. PRODUCTO DE APRENDIZAJE")]
    pr = d.get("producto") or {}
    b.append(("tabla", ["Evidencia", "Criterios de evaluación", "Relación con otras áreas"],
              [[pr.get("evidencia"), pr.get("criterios"), pr.get("relacion")]]))
    b.append(("sec", "IV. SECUENCIA DE LA UNIDAD"))
    b.append(("tabla", ["Semana", "Sesión", "Título / pregunta retadora", "Competencia"],
              [[f"Semana {f.get('semana')}", f"Sesión {i}", f.get("titulo"), f.get("competencia")]
               for i, f in enumerate(d.get("secuencia", []), 1)]))
    b.append(("sec", "V. PROPÓSITOS DE APRENDIZAJE"))
    filas = []
    for p in d.get("propositos", []):
        des = _bul(p.get("desempenos")) + ("\n(Desempeños por verificar con el Programa Curricular)" if p.get("por_verificar") else "")
        filas.append([p.get("competencia"), _bul(p.get("capacidades")), des, p.get("proposito"),
                      p.get("campo_tematico"), p.get("evidencia"), p.get("instrumento")])
    b.append(("tabla", ["Competencia", "Capacidades", "Desempeños", "Propósito", "Campo temático",
                        "Evidencia", "Instrumento"], filas))
    return b + _bloques_transversales(trans, d.get("enfoques") or None)


def bloques_sesion(asig, udata, d, trans):
    dur = sum(int(d.get(k, {}).get("minutos") or 0) for k in ("inicio", "desarrollo", "cierre"))
    b = [("h", f"SESIÓN DE APRENDIZAJE N.° {d.get('numero', '')}: {d.get('titulo', '')}"),
         ("sec", "I. DATOS INFORMATIVOS"),
         ("kv", _kv_datos(asig, [("Unidad", f"{udata.get('numero', '')}: {udata.get('titulo', '')}"),
                                 ("Duración", f"{d.get('duracion') or dur} minutos")]), 2),
         ("sec", "II. PROPÓSITO DE APRENDIZAJE"), ("p", d.get("proposito")),
         ("tabla", ["Competencia", "Capacidades", "Desempeños", "Criterios de evaluación"],
          [[c.get("competencia"), _bul(c.get("capacidades")),
            _bul(c.get("desempenos")) + ("\n(Por verificar)" if c.get("por_verificar") else ""),
            _bul(c.get("criterios"))] for c in d.get("competencias", [])]),
         ("sec", "III. EVALUACIÓN DE LOS APRENDIZAJES"),
         ("tabla", ["Competencia", "Evidencia de aprendizaje", "Instrumento de evaluación"],
          [[c.get("competencia"), c.get("evidencia"), c.get("instrumento")] for c in d.get("competencias", [])])]
    b += _bloques_transversales(trans, d.get("enfoques") or None)
    b.append(("sec", "IV. SECUENCIA DIDÁCTICA"))
    b.append(("tabla", ["Momento", "Actividades", "Tiempo"],
              [[nombre, (d.get(k) or {}).get("actividades"), f"{(d.get(k) or {}).get('minutos', '')} min"]
               for nombre, k in (("INICIO", "inicio"), ("DESARROLLO", "desarrollo"), ("CIERRE", "cierre"))]))
    b += [("sec", "V. MATERIALES Y RECURSOS"), ("ul", d.get("materiales") or []),
          ("sec", "VI. REFLEXIÓN SOBRE EL APRENDIZAJE"), ("ul", d.get("reflexion") or [])]
    return b


def bloques_rubrica(asig, sdata, d):
    return [("h", f"RÚBRICA DE EVALUACIÓN — {sdata.get('titulo', '')}"),
            ("kv", [("Institución educativa", asig.get("institucion")), ("Docente", asig.get("docente")),
                    ("Área", asig.get("area")), ("Grado y sección", f"{asig.get('grado', '')} {asig.get('seccion') or ''}"),
                    ("Competencia(s)", ", ".join(c.get("competencia", "") for c in sdata.get("competencias", [])))], 2),
            ("sec", "CRITERIOS Y NIVELES DE LOGRO"),
            ("tabla", ["Criterio", "AD · Logro destacado", "A · Logro esperado", "B · En proceso", "C · En inicio"],
             [[r.get("criterio"), r.get("AD"), r.get("A"), r.get("B"), r.get("C")] for r in d.get("criterios", [])])]


def bloques_ficha(asig, sdata, d):
    b = [("h", (d.get("titulo") or f"FICHA DE TRABAJO — {sdata.get('titulo', '')}").upper()),
         ("kv", [("Nombres y apellidos", ""), ("Fecha", ""), ("Grado y sección", f"{asig.get('grado', '')} {asig.get('seccion') or ''}"),
                 ("Área", asig.get("area"))], 2),
         ("sec", "INDICACIONES"), ("p", d.get("instrucciones")), ("sec", "ACTIVIDADES")]
    for a in d.get("actividades", []):
        b.append(("p", f"{a.get('numero')}. {s(a.get('consigna'))}"))
        b.append(("p", "Respuesta: ______________________________________________"))
    b += [("salto",), ("sec", "SOLUCIONARIO (SOLO PARA EL DOCENTE)"),
          ("tabla", ["N.°", "Respuesta esperada"], [[a.get("numero"), a.get("respuesta")] for a in d.get("actividades", [])])]
    return b


def bloques_examen(asig, sdata, d):
    def _num(v):
        try:
            return float(str(v).replace(",", "."))
        except ValueError:
            return 0.0
    total = sum(_num(p.get("puntaje")) for p in d.get("preguntas", []))
    b = [("h", (d.get("titulo") or f"EXAMEN — {sdata.get('titulo', '')}").upper()),
         ("kv", [("Nombres y apellidos", ""), ("Fecha", ""), ("Grado y sección", f"{asig.get('grado', '')} {asig.get('seccion') or ''}"),
                 ("Puntaje total", f"{total:g} puntos")], 2),
         ("sec", "INDICACIONES"), ("p", d.get("instrucciones")), ("sec", "PREGUNTAS")]
    for p in d.get("preguntas", []):
        b.append(("sub", f"{p.get('numero')}. {s(p.get('enunciado'))}  ({p.get('puntaje')} pts)"))
        ops = p.get("opciones") or []
        if ops:
            b.append(("ul", ops))
        elif p.get("tipo") == "desarrollo":
            b.append(("p", "______________________________________________________________"))
    b += [("salto",), ("sec", "CLAVES DE RESPUESTA (SOLO PARA EL DOCENTE)"),
          ("tabla", ["N.°", "Respuesta correcta", "Puntaje"],
           [[p.get("numero"), p.get("respuesta"), p.get("puntaje")] for p in d.get("preguntas", [])])]
    return b


# ------------------------------------------------------------------
# WORD
# ------------------------------------------------------------------
def _sombrear(celda, hex_color):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _escribir(celda, texto, negrita=False, color_texto=None, tam=9, centrado=False):
    celda.text = ""
    p = celda.paragraphs[0]
    for i, linea in enumerate(s(texto).split("\n")):
        if i > 0:
            p = celda.add_paragraph()
        run = p.add_run(linea)
        run.bold = negrita
        run.font.size = Pt(tam)
        if color_texto:
            run.font.color.rgb = RGBColor.from_string(color_texto)
        p.paragraph_format.space_after = Pt(1)
        if centrado:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def a_docx(bloques, color=COLOR_DEFECTO, pie: Optional[str] = None, horizontal=False) -> bytes:
    color = (color or COLOR_DEFECTO).lstrip("#")
    claro = _aclarar(color)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    if horizontal:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(1.6))
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    for blq in bloques:
        tipo = blq[0]
        if tipo == "h":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(s(blq[1]))
            r.bold = True
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor.from_string(color)
        elif tipo == "sec":
            t = doc.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            _sombrear(t.rows[0].cells[0], color)
            _escribir(t.rows[0].cells[0], blq[1], negrita=True, color_texto="FFFFFF", tam=10)
        elif tipo == "kv":
            pares, ncols = blq[1], blq[2]
            filas = (len(pares) + ncols - 1) // ncols
            t = doc.add_table(rows=filas, cols=ncols * 2)
            t.style = "Table Grid"
            for idx, (k, v) in enumerate(pares):
                fila, col = divmod(idx, ncols)
                ck, cv = t.rows[fila].cells[col * 2], t.rows[fila].cells[col * 2 + 1]
                _sombrear(ck, claro)
                _escribir(ck, k, negrita=True)
                _escribir(cv, v)
        elif tipo == "tabla":
            enc, filas = blq[1], blq[2]
            t = doc.add_table(rows=1, cols=len(enc))
            t.style = "Table Grid"
            for i, h in enumerate(enc):
                _sombrear(t.rows[0].cells[i], color)
                _escribir(t.rows[0].cells[i], h, negrita=True, color_texto="FFFFFF", centrado=True)
            for fila in filas:
                celdas = t.add_row().cells
                for i in range(len(enc)):
                    _escribir(celdas[i], fila[i] if i < len(fila) else "")
        elif tipo == "p":
            if s(blq[1]):
                doc.add_paragraph(s(blq[1]))
        elif tipo == "sub":
            p = doc.add_paragraph()
            p.add_run(s(blq[1])).bold = True
        elif tipo == "ul":
            for item in blq[1]:
                if s(item):
                    doc.add_paragraph(s(item), style="List Bullet")
        elif tipo == "salto":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if tipo in ("sec", "kv", "tabla"):
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    r = p.add_run("Normativa utilizada: " + (pie or "RM N.° 501-2025-MINEDU") +
                  ". Documento generado con EduPerú IA; revísalo y ajústalo a tu realidad antes de usarlo.")
    r.italic = True
    r.font.size = Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def texto_a_docx(titulo: str, contenido: str, pie: Optional[str] = None) -> bytes:
    """Word sencillo para documentos antiguos (texto con formato Markdown)."""
    doc = Document()
    doc.add_heading(titulo, level=0)
    for linea in contenido.splitlines():
        l = linea.strip().replace("**", "")
        if not l:
            continue
        if l.startswith("### "):
            doc.add_heading(l[4:], level=3)
        elif l.startswith("## "):
            doc.add_heading(l[3:], level=2)
        elif l.startswith("# "):
            doc.add_heading(l[2:], level=1)
        elif l.startswith(("- ", "* ")):
            doc.add_paragraph(l[2:], style="List Bullet")
        else:
            doc.add_paragraph(l)
    if pie:
        doc.add_paragraph().add_run(f"Normativa utilizada: {pie}").italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------
# VISTA PREVIA (HTML)
# ------------------------------------------------------------------
_CSS = """<style>
body{font-family:Arial,Helvetica,sans-serif;font-size:13px;color:#222;margin:0;padding:14px;background:#fff}
h1{text-align:center;font-size:17px;color:#__C__;margin:6px 0 14px}
.sec{background:#__C__;color:#fff;font-weight:bold;padding:6px 10px;margin:14px 0 6px;border-radius:3px;font-size:13px}
table{border-collapse:collapse;width:100%;margin-bottom:8px;table-layout:auto}
th{background:#__C__;color:#fff;padding:5px;border:1px solid #__C__;font-size:12px}
td{border:1px solid #bbb;padding:5px;vertical-align:top;font-size:12px}
td.k{background:#__L__;font-weight:bold;width:16%}
p{margin:4px 0;line-height:1.45} ul{margin:4px 0 8px 18px;padding:0}
.sub{font-weight:bold;margin:8px 0 2px}.salto{border-top:2px dashed #bbb;margin:16px 0}
.pie{font-size:11px;color:#666;font-style:italic;margin-top:14px}
</style>"""


def _h(texto) -> str:
    return _html.escape(s(texto)).replace("\n", "<br>")


def a_html(bloques, color=COLOR_DEFECTO, pie: Optional[str] = None) -> str:
    color = (color or COLOR_DEFECTO).lstrip("#")
    out = [_CSS.replace("__C__", color).replace("__L__", _aclarar(color))]
    for blq in bloques:
        tipo = blq[0]
        if tipo == "h":
            out.append(f"<h1>{_h(blq[1])}</h1>")
        elif tipo == "sec":
            out.append(f"<div class='sec'>{_h(blq[1])}</div>")
        elif tipo == "kv":
            pares, ncols = blq[1], blq[2]
            out.append("<table>")
            for i in range(0, len(pares), ncols):
                out.append("<tr>")
                for k, v in pares[i:i + ncols]:
                    out.append(f"<td class='k'>{_h(k)}</td><td>{_h(v)}</td>")
                out.append("</tr>")
            out.append("</table>")
        elif tipo == "tabla":
            enc, filas = blq[1], blq[2]
            out.append("<table><tr>" + "".join(f"<th>{_h(e)}</th>" for e in enc) + "</tr>")
            for fila in filas:
                out.append("<tr>" + "".join(f"<td>{_h(fila[i] if i < len(fila) else '')}</td>"
                                           for i in range(len(enc))) + "</tr>")
            out.append("</table>")
        elif tipo == "p":
            if s(blq[1]):
                out.append(f"<p>{_h(blq[1])}</p>")
        elif tipo == "sub":
            out.append(f"<div class='sub'>{_h(blq[1])}</div>")
        elif tipo == "ul":
            out.append("<ul>" + "".join(f"<li>{_h(i)}</li>" for i in blq[1] if s(i)) + "</ul>")
        elif tipo == "salto":
            out.append("<div class='salto'></div>")
    out.append(f"<div class='pie'>Normativa utilizada: {_h(pie or 'RM N.° 501-2025-MINEDU')}</div>")
    return "".join(out)
